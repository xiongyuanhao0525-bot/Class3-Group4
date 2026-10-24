# -*- coding: utf-8 -*-
"""回读验证 OOXML 候选：表格与字体修改是否生效。"""
from docx import Document
from docx.oxml.ns import qn

CAND = r'C:\Users\16570\AppData\Local\Doubao\User Data\Default\DoubaoStorage\Aida\doubao-office\ai-edit\2B2EO7HaNFl3HHLAhd50GPpRXoW-1792594847870-fbfwo2yge9g.docx'
doc = Document(CAND)

t = doc.tables[0]
print('rows:', len(t.rows))
grid = t._tbl.find(qn('w:tblGrid'))
ws = [c.get(qn('w:w')) for c in grid.findall(qn('w:gridCol'))]
print('gridCol widths:', ws)
print('hdr4:', t.rows[0].cells[3].text)
th_count = 0
for tr in t._tbl.findall(qn('w:tr')):
    trPr = tr.find(qn('w:trPr'))
    if trPr is not None and trPr.find(qn('w:tblHeader')) is not None:
        th_count += 1
print('tblHeader count:', th_count)

samples = 0
for p in doc.paragraphs:
    try:
        st = (p.style.name or '')
    except Exception:
        st = ''
    if st.lower().startswith('heading') or 'title' in st.lower() or 'caption' in st.lower():
        continue
    for run in p.runs:
        if run.text.strip() and samples < 5:
            rPr = run._r.find(qn('w:rPr'))
            ea = None
            if rPr is not None:
                rf = rPr.find(qn('w:rFonts'))
                if rf is not None:
                    ea = rf.get(qn('w:eastAsia'))
            print('sample run:', repr(run.text[:14]), '| name=', run.font.name, '| size=', run.font.size.pt if run.font.size else None, '| eastAsia=', ea)
            samples += 1
print('sampled:', samples)
