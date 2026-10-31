# -*- coding: utf-8 -*-
"""编辑 OOXML 候选：①表格列宽/表头改短/取消跨页重复表头；②全文正文与表格字体=宋体小四，英文 Times New Roman。"""
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

CAND = Path(r'C:\Users\16570\AppData\Local\Doubao\User Data\Default\DoubaoStorage\Aida\doubao-office\ai-edit\2B2EO7HaNFl3HHLAhd50GPpRXoW-1792594847870-fbfwo2yge9g.docx')

doc = Document(str(CAND))


def set_run_font(run):
    run.font.name = 'Times New Roman'           # ascii + hAnsi
    run.font.size = Pt(12)                       # 小四
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), '宋体')


HEADING_NAMES = ('heading', 'title', 'subtitle', '标题', '题注', 'caption')


def is_heading(p):
    try:
        name = (p.style.name or '').lower()
    except Exception:
        return False
    return any(k in name for k in HEADING_NAMES)


# 1) 表格：列宽 + 表头第4列短文本 + 取消 tblHeader
for table in doc.tables:
    grid = table._tbl.find(qn('w:tblGrid'))
    cols = grid.findall(qn('w:gridCol'))
    widths = ['1200', '1300', '3762', '2600']
    for c, w in zip(cols, widths):
        c.set(qn('w:w'), w)
    hdr_tr = table.rows[0]._tr
    hdr_tcs = hdr_tr.findall(qn('w:tc'))
    if len(hdr_tcs) >= 4:
        # 清空第4列表头文字，写入短文本
        p0 = hdr_tcs[3].findall(qn('w:p'))[0]
        for r in p0.findall(qn('w:r')):
            p0.remove(r)
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = '代码注释标识'
        t.set(qn('xml:space'), 'preserve')
        r.append(t)
        p0.append(r)
    # 移除所有行 tblHeader（不跨页重复表头），保留 cantSplit
    for tr in table._tbl.findall(qn('w:tr')):
        trPr = tr.find(qn('w:trPr'))
        if trPr is not None:
            th = trPr.find(qn('w:tblHeader'))
            if th is not None:
                trPr.remove(th)
    # 表格内所有段落文字：宋体小四 + TNR
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run)

# 2) 正文段落（非标题）：宋体小四 + TNR
changed = 0
for p in doc.paragraphs:
    if is_heading(p):
        continue
    for run in p.runs:
        if run.text.strip():
            set_run_font(run)
            changed += 1

doc.save(str(CAND))
print('font_runs_changed:', changed)
print('tables:', len(doc.tables))
