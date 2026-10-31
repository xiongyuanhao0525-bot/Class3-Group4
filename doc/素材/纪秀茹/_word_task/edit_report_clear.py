# -*- coding: utf-8 -*-
"""增量编辑 doc\\开源软件质量分析报告-DjangoBlog.docx：
①3.1/3.2 冗长段落拆分为编号列表，内容更清晰；
②表格格式修正（列宽/表头短文本/去跨页重复表头）；
③全文字体=宋体小四、英文 Times New Roman。"""
from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
import copy

TARGET = Path(r'C:\Users\16570\Class3-Group4\doc\开源软件质量分析报告-DjangoBlog.docx')
doc = Document(str(TARGET))
paras = doc.paragraphs


def black_run(r_el):
    rPr = r_el.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        r_el.insert(0, rPr)
    color = rPr.find(qn('w:color'))
    if color is None:
        color = OxmlElement('w:color')
        rPr.append(color)
    color.set(qn('w:val'), '000000')


def set_text(p, text):
    runs = p.runs
    if runs:
        runs[0].text = text
        black_run(runs[0]._r)
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        black_run(p.add_run(text)._r)


def insert_after(anchor_p, text):
    new_p_el = OxmlElement('w:p')
    pPr = anchor_p._p.pPr
    if pPr is not None:
        new_p_el.append(copy.deepcopy(pPr))
    anchor_p._p.addnext(new_p_el)
    new_par = Paragraph(new_p_el, anchor_p._parent)
    black_run(new_par.add_run(text)._r)
    return new_par


def find_para(prefix):
    for p in paras:
        if p.text.startswith(prefix):
            return p
    raise RuntimeError('not found: ' + prefix)


# ============ 3.1 代码风格：拆分为结论+正例/反例列表 ============
p31a = find_para('项目整体遵循 PEP8')
set_text(p31a, '项目整体遵循 PEP8 编码风格：统一 4 空格缩进（blog 应用 0 行 Tab）、类名大驼峰、函数与变量小写蛇形、常量全大写，import 分组清晰（标准库→第三方→项目内）。总体代码风格良好，仅有少量历史遗留问题。正例方面：')
pos = [
    '（1）blog/models.py：viewed() 方法三行说清“浏览量+1”，并用 update_fields 显式表达“只想改一列”的意图。',
    '（2）oauth：BaseOauthManager 将 AUTH_URL/TOKEN_URL/API_URL/ICON_NAME 四个类属性集中声明，各平台子类只覆盖差异点。',
    '（3）servermanager：models.py 每个字段均带中文 verbose_name，urls.py 用 app_name 命名空间隔离路由。',
    '（4）djangoblog：apps.py 的 ready() 方法内延迟导入规避循环依赖，hook_constants.py 集中定义钩子常量杜绝魔法字符串。',
]
anchor = p31a
for s in pos:
    anchor = insert_after(anchor, s)

p31b = find_para('反例方面')
set_text(p31b, '反例方面，存在以下需要改进之处：')
neg = [
    '（1）超长单行：blog_tags.py 第 872 行单行 141 字符；views.py 第 217/220 行 126 字符，且“#commentlist-container”锚点硬编码两遍。',
    '（2）字符串格式化混用：同一模块内 f-string、.format()、% 三种风格并存（models.py:316、blog_tags.py:553、ping_baidu.py:71）。',
    '（3）历史拼写错误（不建议改，涉及模板键名与外部引用）：commment_list、sidabar_links、ElaspedTimeDocumentManager。',
    '（4）注释与行宽（oauth）：中英混杂的魔法注释（如“facebook 的 token 过长”）、90+ 字符的 URL 常量行，且项目无统一行宽配置文件。',
    '（5）其他（servermanager/djangoblog）：robot.py 续行多缩进 14 空格（flake8 E127）并硬编码开发 token；settings.py 有 25 行被注释掉的旧配置；whoosh_cn_backend.py 以非 raw 字符串书写正则 \\d。',
    '结论：代码风格总体遵循 PEP8，个别历史遗留问题已登记入改进计划。',
]
anchor = p31b
for s in neg:
    anchor = insert_after(anchor, s)

# ============ 3.2 代码设计：拆分为结论+优点/耦合点列表 ============
p32a = find_para('项目总体上遵循')
set_text(p32a, '项目总体上遵循模块化、信息隐藏、高内聚、低耦合四项设计原则，可维护性较好，具体评价如下：')
des = [
    '（1）模块化：blog 应用按“路由→视图→模型/模板标签/表单+横切（中间件、搜索、后台、运维命令）”分层清晰；oauth 应用将第三方平台 HTTP 交互细节独立到 oauthmanager.py 策略层，视图层只编排流程；servermanager 的 robot.py 只做消息路由与状态机，api/ 目录按单一职责拆分；djangoblog 分 5 层职责（框架配置/横切基础设施/事件/后台/插件框架）。',
    '（2）信息隐藏：djangoblog/base_views.py 把 CSRF、登录校验、禁缓存等横切保护收敛到 dispatch 层；blog 的 BlogSettings 单例把“仅一份配置”约束落在数据层 clean() 与后台三个权限钩子上；vite_tags 的 load_manifest() 收敛 manifest 读取与兜底逻辑。',
]
anchor = p32a
for s in des:
    anchor = insert_after(anchor, s)

p32b = find_para('高内聚：blog 的 ElaspedTimeDocumentManager')
set_text(p32b, '（3）高内聚与低耦合：blog 的 ElaspedTimeDocumentManager 把建索引/删索引/写记录收在同一个类；oauth 的 BaseOauthManager 采用策略模式+模板方法，5 个平台子类各约 60~90 行只实现差异部分，对修改关闭、对扩展开放；comments 的 Comment 模型以自引用外键实现楼中楼，CommentReaction 用 unique_together 保证同一用户对同一评论的同一表情唯一。低耦合方面，comments 视图继承 AuthenticatedFormView 复用登录与表单流程；oauth 的 ProxyManagerMixin 利用 Python MRO 组合注入代理能力，零 if-else。')
anchor = insert_after(p32b, '不足与可点名的耦合点如下：')
couple = [
    '（1）blog：ES 文档字段与 Article 模型字段需手工同步（隐性契约）；load_sidebar 一个标签直接查询文章/分类/标签等 8 类数据；ModelAdmin 定义与注册分离（blog/admin.py 与 djangoblog/admin_site.py 两处）。',
    '（2）oauth：5 个子类构造函数前 4 行完全重复（复制粘贴 5 次）；authorize 视图约 100 行承担 8 件事。',
    '（3）servermanager：BlogApi 直接 import blog ORM，双向耦合。',
    '（4）djangoblog：blog_signals 事件中心单函数约 110 行集中分发多 app 模型。',
]
for s in couple:
    anchor = insert_after(anchor, s)

# ============ 表格格式修正 ============
for table in doc.tables:
    grid = table._tbl.find(qn('w:tblGrid'))
    cols = grid.findall(qn('w:gridCol'))
    for c, w in zip(cols, ['1200', '1300', '3762', '2600']):
        c.set(qn('w:w'), w)
    hdr_tr = table.rows[0]._tr
    hdr_tcs = hdr_tr.findall(qn('w:tc'))
    if len(hdr_tcs) >= 4:
        p0 = hdr_tcs[3].findall(qn('w:p'))[0]
        for r in p0.findall(qn('w:r')):
            p0.remove(r)
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = '代码注释标识'
        t.set(qn('xml:space'), 'preserve')
        r.append(t)
        p0.append(r)
    for tr in table._tbl.findall(qn('w:tr')):
        trPr = tr.find(qn('w:trPr'))
        if trPr is not None:
            th = trPr.find(qn('w:tblHeader'))
            if th is not None:
                trPr.remove(th)


def set_run_font(run):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), '宋体')


HEADING_NAMES = ('heading', 'title', 'subtitle', '标题', '题注', 'caption')


def is_heading(p):
    try:
        name = (p.style.name or '').lower()
    except Exception:
        return False
    return any(k in name for k in HEADING_NAMES)


for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run)

changed = 0
for p in doc.paragraphs:
    if is_heading(p):
        continue
    for run in p.runs:
        if run.text.strip():
            set_run_font(run)
            changed += 1

doc.save(str(TARGET))
print('saved. font runs changed:', changed, '| paragraphs:', len(doc.paragraphs), '| tables:', len(doc.tables))
